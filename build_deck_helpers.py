import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

# -----------------------------------------------------------------------------
# PALETTE DEFINITIONS (Sahara Obsidian Cyber-Tactical Theme)
# -----------------------------------------------------------------------------
BG_COLOR        = RGBColor(11, 15, 25)      # #0B0F19 - Midnight Obsidian
SURFACE_COLOR   = RGBColor(19, 27, 46)      # #131B2E - Dark Slate Card
SURFACE_ALT     = RGBColor(27, 38, 64)      # #1B2640 - Elevated Card
BORDER_COLOR    = RGBColor(40, 56, 84)      # #283854 - Subtle Accent Border
BORDER_ACTIVE   = RGBColor(245, 158, 11)    # #F59E0B - Sahara Amber Border

TEXT_TITLE      = RGBColor(248, 250, 252)   # #F8FAFC - Pure Off-White
TEXT_BODY       = RGBColor(203, 213, 225)   # #CBD5E1 - Light Slate
TEXT_MUTED      = RGBColor(148, 163, 184)   # #94A3B8 - Slate Muted

ACCENT_GOLD     = RGBColor(245, 158, 11)    # #F59E0B - Sahara Amber / Gold
ACCENT_CYAN     = RGBColor(56, 189, 248)    # #38BDF8 - Tactical Cyan
ACCENT_GREEN    = RGBColor(16, 185, 129)    # #10B981 - Emerald Verified
ACCENT_RED      = RGBColor(239, 68, 68)     # #EF4444 - Crimson Distress / SOS
ACCENT_PURPLE   = RGBColor(168, 85, 247)    # #A855F7 - Cryptographic Violet

FONT_HEADING = "Segoe UI"
FONT_BODY    = "Segoe UI"

# -----------------------------------------------------------------------------
# INITIALIZE PRESENTATION (16:9 Widescreen)
# -----------------------------------------------------------------------------
prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout     = prs.slide_layouts[6]

def set_background(slide):
    """Draws a full-bleed dark obsidian background rectangle."""
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG_COLOR
    bg.line.fill.background()
    return bg

def add_header(slide, tag_text, title_text, subtitle_text=None, tag_color=ACCENT_GOLD):
    """Adds a standardized keynotes header with a category tag, title, and subtitle."""
    # Tag Pill / Badge
    tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.35), Inches(8.0), Inches(0.3))
    tf_tag = tag_box.text_frame
    tf_tag.word_wrap = True
    tf_tag.margin_left = tf_tag.margin_top = tf_tag.margin_right = tf_tag.margin_bottom = 0
    p_tag = tf_tag.paragraphs[0]
    p_tag.text = tag_text.upper()
    p_tag.font.name = FONT_HEADING
    p_tag.font.size = Pt(10)
    p_tag.font.bold = True
    p_tag.font.color.rgb = tag_color

    # Slide Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.62), Inches(11.5), Inches(0.55))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    tf_title.margin_left = tf_title.margin_top = tf_title.margin_right = tf_title.margin_bottom = 0
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.name = FONT_HEADING
    p_title.font.size = Pt(22)
    p_title.font.bold = True
    p_title.font.color.rgb = TEXT_TITLE

    # Slide Subtitle
    if subtitle_text:
        sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.18), Inches(11.5), Inches(0.35))
        tf_sub = sub_box.text_frame
        tf_sub.word_wrap = True
        tf_sub.margin_left = tf_sub.margin_top = tf_sub.margin_right = tf_sub.margin_bottom = 0
        p_sub = tf_sub.paragraphs[0]
        p_sub.text = subtitle_text
        p_sub.font.name = FONT_BODY
        p_sub.font.size = Pt(11)
        p_sub.font.color.rgb = TEXT_MUTED

def add_card(slide, left, top, width, height, bg_color=SURFACE_COLOR, border_color=BORDER_COLOR):
    """Adds a dark card container shape with subtle border."""
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    card.line.color.rgb = border_color
    card.line.width = Pt(1)
    return card

def add_phone_mockup(slide, img_path, left, top, max_height=Inches(5.4)):
    """Embeds an actual screenshot with correct aspect ratio and dark bezels."""
    if not os.path.exists(img_path):
        return None
    try:
        im = Image.open(img_path)
        w, h = im.size
        target_h = max_height
        target_w = target_h * (w / h)

        # Phone frame outline / shadow backplate
        frame_pad = Inches(0.08)
        frame = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            left - frame_pad,
            top - frame_pad,
            target_w + (frame_pad * 2),
            target_h + (frame_pad * 2)
        )
        frame.fill.solid()
        frame.fill.fore_color.rgb = RGBColor(15, 20, 32)
        frame.line.color.rgb = RGBColor(60, 78, 110)
        frame.line.width = Pt(1.5)

        pic = slide.shapes.add_picture(img_path, left, top, width=target_w, height=target_h)
        return pic
    except Exception as e:
        print(f"Error adding picture {img_path}: {e}")
        return None

def add_stat_box(slide, left, top, width, height, big_num, label_text, accent_color=ACCENT_GOLD):
    """Adds a stat metric chip with large number and concise label."""
    card = add_card(slide, left, top, width, height, bg_color=SURFACE_ALT, border_color=BORDER_COLOR)
    
    tb = slide.shapes.add_textbox(left + Inches(0.15), top + Inches(0.08), width - Inches(0.3), height - Inches(0.16))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    
    p1 = tf.paragraphs[0]
    p1.text = big_num
    p1.font.name = FONT_HEADING
    p1.font.size = Pt(22)
    p1.font.bold = True
    p1.font.color.rgb = accent_color
    
    p2 = tf.add_paragraph()
    p2.text = label_text
    p2.font.name = FONT_BODY
    p2.font.size = Pt(9.5)
    p2.font.color.rgb = TEXT_MUTED
    return card

print("Helper definitions loaded successfully.")
